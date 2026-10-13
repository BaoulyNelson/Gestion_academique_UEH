from decimal import Decimal, InvalidOperation
from zipfile import BadZipFile

from django.core.exceptions import ValidationError
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

ENTETES = ("code_etudiant", "evaluation", "note", "coefficient")
MAXIMUM_LIGNES = 5000


def lire_evaluations_excel(fichier):
    try:
        classeur = load_workbook(fichier, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as erreur:
        raise ValidationError("Le fichier Excel est illisible ou invalide.") from erreur

    erreurs = []
    evaluations = []
    try:
        feuille = classeur.active
        lignes = feuille.iter_rows(values_only=True)
        entetes = next(lignes, None)
        entetes_normalisees = tuple(
            str(valeur).strip().lower() if valeur is not None else ""
            for valeur in (entetes or ())
        )
        if entetes_normalisees[: len(ENTETES)] != ENTETES:
            raise ValidationError(
                "Les colonnes requises sont: code_etudiant, evaluation, note, coefficient."
            )

        for numero_ligne, valeurs in enumerate(lignes, start=2):
            if not any(
                valeur is not None and str(valeur).strip() for valeur in valeurs
            ):
                continue
            if len(evaluations) >= MAXIMUM_LIGNES:
                erreurs.append(
                    f"Le fichier depasse la limite de {MAXIMUM_LIGNES} lignes."
                )
                break

            code, libelle, note_brute, coefficient_brut = (
                valeurs[index] if index < len(valeurs) else None
                for index in range(len(ENTETES))
            )
            code = str(code).strip() if code is not None else ""
            libelle = str(libelle).strip() if libelle is not None else ""
            if not code or not libelle or len(libelle) > 120:
                erreurs.append(
                    f"Ligne {numero_ligne}: le code et le nom de l'evaluation sont requis (120 caracteres maximum)."
                )
                continue

            try:
                note = Decimal(str(note_brute).replace(",", "."))
                coefficient = Decimal(str(coefficient_brut).replace(",", "."))
                if not note.is_finite() or not Decimal("0") <= note <= Decimal("100"):
                    raise InvalidOperation
                if not coefficient.is_finite() or coefficient <= 0:
                    raise InvalidOperation
            except (InvalidOperation, TypeError, ValueError):
                erreurs.append(
                    f"Ligne {numero_ligne}: la note doit etre entre 0 et 100 et le coefficient superieur a 0."
                )
                continue

            evaluations.append(
                {
                    "ligne": numero_ligne,
                    "code_etudiant": code,
                    "libelle": libelle,
                    "note": note,
                    "coefficient": coefficient,
                }
            )
    finally:
        classeur.close()

    if not evaluations:
        erreurs.append("Le fichier ne contient aucune evaluation exploitable.")
    if erreurs:
        raise ValidationError(erreurs)
    return evaluations
